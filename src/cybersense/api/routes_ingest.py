"""Ingest routes — POST /documents (upload), GET/DELETE /documents, POST /ingest/{source}."""

from __future__ import annotations

import hashlib
import tempfile
import uuid
from pathlib import Path

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from cybersense.api.deps import get_embedder, get_vector_store, verify_api_key
from cybersense.config import Settings, get_settings
from cybersense.db.session import get_session
from cybersense.db.tables import DocumentRecord, IngestionRun
from cybersense.generation.guardrails import validate_upload_filename
from cybersense.ingestion.pipeline import IngestionPipeline
from cybersense.models import DocType, DocumentListItem, IngestionRunStatus
from cybersense.retrieval.embeddings import EmbeddingModel
from cybersense.retrieval.vector_store import VectorStore

router = APIRouter(tags=["ingest"])


@router.post("/documents", dependencies=[Depends(verify_api_key)])
async def upload_document(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_session),
    settings: Settings = Depends(get_settings),
    embedder: EmbeddingModel = Depends(get_embedder),
    vector_store: VectorStore = Depends(get_vector_store),
):
    """Upload a PDF/DOCX/MD/TXT/HTML file and ingest it."""
    filename = file.filename or "upload"
    if not validate_upload_filename(filename, settings.api.allowed_extensions):
        raise HTTPException(status_code=400, detail=f"Disallowed file extension. Permitted: {settings.api.allowed_extensions}")

    content = await file.read()
    if len(content) > settings.api.upload_max_mb * 1024 * 1024:
        raise HTTPException(status_code=413, detail=f"File exceeds {settings.api.upload_max_mb} MB limit")

    suffix = Path(filename).suffix.lower()
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        tmp.write(content)
        tmp_path = Path(tmp.name)

    try:
        from cybersense.ingestion.pdf import PDFLoader
        from cybersense.ingestion.markdown import MarkdownLoader
        from cybersense.ingestion.html import HTMLLoader

        if suffix == ".pdf":
            loader = PDFLoader(source_label="upload", doc_type=DocType.upload)
            docs = list(loader.load(paths=[tmp_path]))
        elif suffix in (".md", ".markdown"):
            loader = MarkdownLoader(source_label="upload", doc_type=DocType.upload)
            docs = list(loader.load(paths=[tmp_path]))
        elif suffix in (".html", ".htm"):
            loader = HTMLLoader(source_label="upload", doc_type=DocType.upload)
            docs = list(loader.load(paths=[tmp_path]))
        elif suffix == ".txt":
            text = content.decode("utf-8", errors="replace")
            from cybersense.models import Document
            docs = [Document(
                doc_id=f"upload:{hashlib.sha256(content).hexdigest()[:12]}",
                source="upload", title=filename, doc_type=DocType.upload, text=text,
            )]
        elif suffix == ".docx":
            from docx import Document as DocxDocument
            docx = DocxDocument(str(tmp_path))
            text = "\n".join(p.text for p in docx.paragraphs)
            from cybersense.models import Document
            docs = [Document(
                doc_id=f"upload:{hashlib.sha256(content).hexdigest()[:12]}",
                source="upload", title=filename, doc_type=DocType.upload, text=text,
            )]
        else:
            raise HTTPException(status_code=400, detail="Unsupported file type")

        pipeline = IngestionPipeline(db, vector_store, embedder, settings)
        run = await pipeline.run("upload", iter(docs))
        return {"status": "ok", "added": run.added, "updated": run.updated, "doc_ids": [d.doc_id for d in docs]}
    finally:
        tmp_path.unlink(missing_ok=True)


@router.get("/documents", response_model=list[DocumentListItem], dependencies=[Depends(verify_api_key)])
async def list_documents(
    source: str | None = None,
    doc_type: str | None = None,
    db: AsyncSession = Depends(get_session),
):
    stmt = select(DocumentRecord)
    if source:
        stmt = stmt.where(DocumentRecord.source == source)
    if doc_type:
        stmt = stmt.where(DocumentRecord.doc_type == doc_type)
    result = await db.execute(stmt)
    return [DocumentListItem(**{c.name: getattr(r, c.name) for c in DocumentRecord.__table__.columns}) for r in result.scalars()]


@router.delete("/documents/{doc_id}", dependencies=[Depends(verify_api_key)])
async def delete_document(
    doc_id: str,
    db: AsyncSession = Depends(get_session),
    vector_store: VectorStore = Depends(get_vector_store),
):
    stmt = select(DocumentRecord).where(DocumentRecord.doc_id == doc_id)
    result = await db.execute(stmt)
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=404, detail="Document not found")

    await vector_store.delete_document(doc_id)
    await db.delete(record)
    await db.commit()
    return {"deleted": doc_id}


@router.post("/ingest/{source}", dependencies=[Depends(verify_api_key)])
async def trigger_ingest(
    source: str,
    since: str | None = None,
    limit: int | None = None,
    db: AsyncSession = Depends(get_session),
    settings: Settings = Depends(get_settings),
    embedder: EmbeddingModel = Depends(get_embedder),
    vector_store: VectorStore = Depends(get_vector_store),
):
    """Trigger ingestion of a named source (admin endpoint)."""
    from cybersense.scripts_lib import get_loader_for_source
    try:
        loader = get_loader_for_source(source, settings)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))

    docs = loader.load(since=since, limit=limit)
    pipeline = IngestionPipeline(db, vector_store, embedder, settings)
    run = await pipeline.run(source, docs)
    return IngestionRunStatus(
        id=run.id, source=run.source, started_at=run.started_at,
        finished_at=run.finished_at, seen=run.seen, added=run.added,
        updated=run.updated, skipped=run.skipped, failed=run.failed,
        status=run.status, error=run.error,
    )


@router.get("/ingest/runs", response_model=list[IngestionRunStatus], dependencies=[Depends(verify_api_key)])
async def list_ingest_runs(db: AsyncSession = Depends(get_session)):
    result = await db.execute(select(IngestionRun).order_by(IngestionRun.started_at.desc()).limit(50))
    return [
        IngestionRunStatus(
            id=r.id, source=r.source, started_at=r.started_at,
            finished_at=r.finished_at, seen=r.seen, added=r.added,
            updated=r.updated, skipped=r.skipped, failed=r.failed,
            status=r.status, error=r.error,
        )
        for r in result.scalars()
    ]
