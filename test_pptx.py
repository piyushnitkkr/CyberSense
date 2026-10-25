import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent / "src"))

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Add title slide
slide_layout = prs.slide_layouts[6] # blank
slide = prs.slides.add_slide(slide_layout)

print("python-pptx works successfully!")
